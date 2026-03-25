"""Document parsing utilities for DOCX, PDF, Excel, CSV."""

import csv
import io
import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


def parse_docx(file_path: str) -> str:
    """Extract text from a DOCX file."""
    from docx import Document

    doc = Document(file_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            tables_text.append(" | ".join(cells))

    full_text = "\n".join(paragraphs)
    if tables_text:
        full_text += "\n\n--- Таблицы ---\n" + "\n".join(tables_text)
    return full_text


def parse_pdf_text(file_path: str) -> str:
    """Extract text from a PDF using PyPDF2 (text-based PDFs)."""
    from PyPDF2 import PdfReader

    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            text_parts.append(text)
    return "\n\n".join(text_parts)


def parse_excel(file_path: str, sheet_name: str | None = None) -> list[dict[str, Any]]:
    """Parse an Excel file into a list of row dicts."""
    import pandas as pd

    df = pd.read_excel(file_path, sheet_name=sheet_name or 0)
    df = df.fillna("")
    return df.to_dict(orient="records")


def parse_csv(file_path: str, encoding: str = "utf-8") -> list[dict[str, Any]]:
    """Parse a CSV file into a list of row dicts."""
    with open(file_path, "r", encoding=encoding) as f:
        reader = csv.DictReader(f)
        return list(reader)


def detect_and_parse(file_path: str) -> str:
    """Detect file type and extract text."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".docx":
        return parse_docx(file_path)
    elif ext == ".pdf":
        return parse_pdf_text(file_path)
    elif ext in (".xlsx", ".xls"):
        rows = parse_excel(file_path)
        return "\n".join(str(row) for row in rows)
    elif ext == ".csv":
        rows = parse_csv(file_path)
        return "\n".join(str(row) for row in rows)
    elif ext in (".txt", ".text"):
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    else:
        logger.warning("Unsupported file type: %s", ext)
        return ""


def extract_text_with_structure(file_path: str) -> dict[str, Any]:
    """Extract text with structural metadata."""
    path = Path(file_path)
    ext = path.suffix.lower()

    result: dict[str, Any] = {
        "file_name": path.name,
        "file_type": ext,
        "text": "",
        "sections": [],
    }

    text = detect_and_parse(file_path)
    result["text"] = text

    # Simple section detection by numbered headers
    lines = text.split("\n")
    current_section = None
    section_lines: list[str] = []

    for line in lines:
        stripped = line.strip()
        # Detect sections like "1.", "1.1.", "Раздел 1" etc.
        import re
        if re.match(r"^\d+\.(\d+\.)*\s+", stripped) or re.match(r"^(Раздел|Статья|Глава)\s+\d+", stripped):
            if current_section:
                result["sections"].append({
                    "title": current_section,
                    "content": "\n".join(section_lines),
                })
            current_section = stripped
            section_lines = []
        elif current_section:
            section_lines.append(stripped)

    if current_section:
        result["sections"].append({
            "title": current_section,
            "content": "\n".join(section_lines),
        })

    return result
